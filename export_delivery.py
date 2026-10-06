from decimal import Decimal
from pathlib import Path
import sys

from openpyxl import load_workbook

from app.excel_exporter import (
    ExcelExportError,
    export_delivery,
)
from app.models import DeliveryData, EXCEL_HEADERS
from app.parser import PdfParseError, parse_delivery
from app.pdf_reader import PdfReadError, read_pdf
from app.validator import (
    ValidationError,
    validate_delivery,
)


def verify_saved_excel(
    path: Path,
    delivery: DeliveryData,
) -> None:
    workbook = load_workbook(path)

    try:
        sheet = workbook.active

        headers = tuple(
            cell.value
            for cell in sheet[1]
        )

        if headers != EXCEL_HEADERS:
            raise ValueError(
                "Saved Excel headings do not match."
            )

        if sheet.max_row - 1 != len(delivery.records):
            raise ValueError(
                "Saved Excel row count does not match."
            )

        if sheet.max_column != len(EXCEL_HEADERS):
            raise ValueError(
                "Saved Excel column count does not match."
            )

        if sheet.protection.sheet:
            raise ValueError(
                "The worksheet must not be protected."
            )

        if workbook.security.lockStructure:
            raise ValueError(
                "The workbook structure must not be locked."
            )

        total_meters = Decimal("0")

        for row, record in zip(
            sheet.iter_rows(min_row=2),
            delivery.records,
            strict=True,
        ):
            delivery_cell = row[0]

            if (
                delivery_cell.value != record.delivery_no
                or delivery_cell.data_type != "s"
                or delivery_cell.number_format != "@"
            ):
                raise ValueError(
                    f"Row {delivery_cell.row}: "
                    "Delivery No was not preserved as text."
                )

            expected_identifiers = (
                record.packing_list_no,
                record.po_number,
                record.rm_number,
                record.carton_number,
            )

            actual_identifiers = tuple(
                str(cell.value)
                for cell in row[1:5]
            )

            if actual_identifiers != expected_identifiers:
                raise ValueError(
                    f"Row {delivery_cell.row}: "
                    "An identifier changed during export."
                )

            if row[5].data_type != "n":
                raise ValueError(
                    f"Row {delivery_cell.row}: "
                    "Mts/Box must be numeric."
                )

            saved_quantity = Decimal(str(row[5].value))

            if saved_quantity != record.mts_per_box:
                raise ValueError(
                    f"Row {delivery_cell.row}: "
                    "Mts/Box changed during export."
                )

            total_meters += saved_quantity

        if total_meters != delivery.total_meters:
            raise ValueError(
                "Saved Excel quantity total does not match."
            )

        if any(
            cell.protection.locked
            for row in sheet.iter_rows()
            for cell in row
        ):
            raise ValueError(
                "Populated cells must remain unlocked."
            )

    finally:
        workbook.close()


def main() -> int:
    if len(sys.argv) != 3:
        print(
            'Usage: python export_delivery.py '
            '"samples/document.pdf" '
            '"output/result.xlsx"'
        )
        return 1

    pdf_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])

    try:
        pages = read_pdf(pdf_path)
        delivery = parse_delivery(pages)
        validate_delivery(delivery)

        saved_path = export_delivery(
            delivery,
            output_path,
        )

    except (
        PdfReadError,
        PdfParseError,
        ValidationError,
        ExcelExportError,
    ) as exc:
        print(f"\nERROR:\n{exc}")
        return 1

    try:
        verify_saved_excel(
            saved_path,
            delivery,
        )

    except Exception as exc:
        print(
            "\nThe file was saved, but verification failed."
        )
        print(f"File: {saved_path.resolve()}")
        print(f"Details: {exc}")
        return 1

    print("\nEXCEL EXPORT SUCCESSFUL")
    print("-" * 60)
    print(f"Delivery No: {delivery.delivery_no}")
    print(
        f"Packing Lists: {len(delivery.packing_lists)}"
    )
    print(f"Carton rows: {len(delivery.records)}")
    print(
        f"Total Mts/Box: {delivery.total_meters:,}"
    )

    print("\nVerification:")
    print("  All row values: PASS")
    print("  Delivery No text and leading zeros: PASS")
    print("  Numeric Mts/Box values: PASS")
    print("  Workbook and worksheet editability: PASS")

    print(f"\nSaved file:\n{saved_path.resolve()}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())