from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
    Protection,
)

from app.batch_service import BatchData, validate_batch
from app.excel_exporter import (
    ExcelExportError,
    identifier_to_excel,
    quantity_to_excel,
)
from app.models import EXCEL_HEADERS


def export_batch(
    batch: BatchData,
    output_path: str | Path,
) -> Path:
    validate_batch(batch)

    path = Path(output_path)

    if path.suffix.lower() != ".xlsx":
        raise ExcelExportError(
            "The output filename must end with .xlsx."
        )

    if len(batch.records) > 1_048_575:
        raise ExcelExportError(
            "There are too many rows for one Excel worksheet."
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"

    sheet.append(list(EXCEL_HEADERS))

    for record in batch.records:
        sheet.append(
            [
                record.delivery_no,
                record.packing_list_no,
                identifier_to_excel(record.po_number),
                identifier_to_excel(record.rm_number),
                identifier_to_excel(record.carton_number),
                quantity_to_excel(record.mts_per_box),
            ]
        )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="18324D",
    )

    for cell in sheet[1]:
        cell.font = Font(
            name="Calibri",
            size=11,
            bold=True,
            color="FFFFFF",
        )
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="left",
            vertical="center",
        )

    sheet.row_dimensions[1].height = 26

    for row in sheet.iter_rows(
        min_row=2,
        max_row=sheet.max_row,
        min_col=1,
        max_col=6,
    ):
        for cell in row:
            cell.font = Font(
                name="Calibri",
                size=11,
            )
            cell.alignment = Alignment(
                vertical="center",
                horizontal=(
                    "left"
                    if isinstance(cell.value, str)
                    else "right"
                ),
            )
            cell.number_format = (
                "@"
                if isinstance(cell.value, str)
                else "0"
            )

        # Preserve the source delivery identifier exactly.
        row[0].number_format = "@"

        row[5].number_format = (
            "#,##0"
            if isinstance(row[5].value, int)
            else "#,##0.###############"
        )

    for row in sheet.iter_rows():
        for cell in row:
            cell.protection = Protection(
                locked=False,
                hidden=False,
            )

    sheet.protection.sheet = False

    for column, width in {
        "A": 18,
        "B": 21,
        "C": 20,
        "D": 20,
        "E": 18,
        "F": 18,
    }.items():
        sheet.column_dimensions[column].width = width

    sheet.freeze_panes = "A2"

    # Filter dropdowns for all six columns.
    sheet.auto_filter.ref = (
        f"A1:F{sheet.max_row}"
    )

    created = False

    try:
        path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        # Exclusive creation prevents accidental overwriting.
        with path.open("xb") as stream:
            created = True
            workbook.save(stream)

    except FileExistsError as exc:
        raise ExcelExportError(
            "This Excel file already exists. "
            "Choose another filename."
        ) from exc

    except Exception as exc:
        if created:
            try:
                path.unlink(missing_ok=True)
            except OSError:
                pass

        raise ExcelExportError(
            f"Could not save the combined Excel file: {exc}"
        ) from exc

    finally:
        workbook.close()

    return path