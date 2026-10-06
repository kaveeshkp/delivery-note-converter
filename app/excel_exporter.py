from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
    Protection,
)
from openpyxl.utils import get_column_letter

from app.models import DeliveryData, EXCEL_HEADERS
from app.validator import validate_delivery


class ExcelExportError(Exception):
    """Raised when the Excel file cannot be saved."""


def identifier_to_excel(value: str) -> str | int:
    """
    Preserve identifiers with leading zeros or more
    than 15 digits as text.

    Ordinary numeric identifiers are exported as
    numbers to match the approved sample.
    """
    if not value.isascii() or not value.isdigit():
        return value

    if len(value) > 1 and value.startswith("0"):
        return value

    if len(value) > 15:
        return value

    return int(value)


def quantity_to_excel(value: Decimal) -> int | float:
    if value == value.to_integral_value():
        return int(value)

    return float(value)


def export_delivery(
    delivery: DeliveryData,
    output_path: str | Path,
) -> Path:
    # Never export data that has failed validation.
    validate_delivery(delivery)

    path = Path(output_path)

    if path.suffix.lower() != ".xlsx":
        raise ExcelExportError(
            "The output filename must end with .xlsx."
        )

    # Avoid overwriting an existing file silently.
    # The desktop Save As workflow will handle
    # overwrite confirmation when it is implemented.
    if path.exists():
        raise ExcelExportError(
            f"The output file already exists: {path}\n"
            "Choose a different filename."
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"

    sheet.append(list(EXCEL_HEADERS))

    for record in delivery.records:
        sheet.append(
            [
                record.delivery_no,
                record.packing_list_no,
                identifier_to_excel(record.po_number),
                identifier_to_excel(record.rm_number),
                identifier_to_excel(record.carton_number),
                quantity_to_excel(record.mts_per_box),
            ]
        )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="18324D",
    )

    for cell in sheet[1]:
        cell.font = Font(
            name="Calibri",
            size=11,
            bold=True,
            color="FFFFFF",
        )
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="left",
            vertical="center",
        )

    sheet.row_dimensions[1].height = 26

    for row in sheet.iter_rows(
        min_row=2,
        max_row=sheet.max_row,
        min_col=1,
        max_col=6,
    ):
        for cell in row:
            cell.font = Font(
                name="Calibri",
                size=11,
            )
            cell.alignment = Alignment(
                vertical="center",
                horizontal=(
                    "left"
                    if isinstance(cell.value, str)
                    else "right"
                ),
            )

            if isinstance(cell.value, str):
                cell.number_format = "@"
            else:
                cell.number_format = "0"

        # Delivery No is already a string.
        # Text format preserves its initial presentation.
        row[0].number_format = "@"

        # Show whole quantities without decimal places,
        # and retain decimal detail when present.
        quantity = row[5].value

        row[5].number_format = (
            "#,##0"
            if isinstance(quantity, int)
            else "#,##0.###############"
        )

    # Explicitly keep all populated cells unlocked.
    for row in sheet.iter_rows():
        for cell in row:
            cell.protection = Protection(
                locked=False,
                hidden=False,
            )

    sheet.protection.sheet = False

    column_widths = {
        "A": 18,
        "B": 21,
        "C": 20,
        "D": 20,
        "E": 18,
        "F": 18,
    }

    for column, width in column_widths.items():
        sheet.column_dimensions[column].width = width

    # Keep headings visible when scrolling.
    sheet.freeze_panes = "A2"

    last_column = get_column_letter(len(EXCEL_HEADERS))
    sheet.auto_filter.ref = (
        f"A1:{last_column}{sheet.max_row}"
    )

    try:
        path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )
        workbook.save(path)

    except PermissionError as exc:
        raise ExcelExportError(
            "Cannot save the Excel file. "
            "Check folder permissions and whether "
            "the file is open in another application."
        ) from exc

    except OSError as exc:
        raise ExcelExportError(
            f"Could not save the Excel file: {exc}"
        ) from exc

    finally:
        workbook.close()

    return path